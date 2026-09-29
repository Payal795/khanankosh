# table extractor we have to use something which provide boundary boxes
"""
heading_normalizer.py

pdf2docx preserves visual formatting (bold, font size) but does NOT tag
headings with Word's built-in Heading N paragraph styles, and it often
merges a heading line and the paragraph that follows it into a single
<w:p> separated by a <w:br/>. That breaks downstream topic segmentation.

Additionally, multi-page continuous tables extracted by pdf2docx are split
into separate <w:tbl> elements on each page.

This module:
  1. Uses PyMuPDF (fitz) or DOCX structural analysis to detect continuous
     multi-page tables across page breaks.
  2. Joins continuous tables at the XML level in python-docx, preserving all
     formatting and dropping repeated headers on page breaks.
  3. Scans the document to find body-text font size (mode of non-bold sizes).
  4. Finds distinct (bold, size) combinations larger than body size and
     ranks them into Heading 1..N.
  5. Splits any paragraph containing a <w:br/> between a heading run and body.
  6. Applies built-in Heading 1..N styles to heading paragraphs.

Usage:
    from heading_normalizer import normalize_headings
    normalize_headings("converted.docx", "converted_normalized.docx", pdf_path="original.pdf")
"""

import re
from collections import Counter
from dataclasses import dataclass
from typing import List, Dict, Any, Optional, Tuple

import docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

try:
    import pymupdf  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False

MAX_HEADING_LEVEL = 4          # Cap Heading1..Heading4
SIZE_MARGIN = 1.03             # Heading must be at least 3% larger than body text


@dataclass
class RunInfo:
    text: str
    bold: bool
    italic: bool
    size_pt: Optional[float]
    font_name: Optional[str]
    color_rgb: Optional[str]


# -----------------------------------------------------------------------------
# String & Table Matching Helpers
# -----------------------------------------------------------------------------

def _clean_cell(cell: Any) -> str:
    if cell is None:
        return ""
    text = str(cell)
    return re.sub(r"\s+", " ", text).strip()


def _normalize_text(text: str) -> str:
    cleaned = _clean_cell(text).lower()
    return re.sub(r"[^a-z0-9]+", "", cleaned)


def _headers_similar(headers1: List[str], headers2: List[str]) -> bool:
    if not headers1 or not headers2:
        return False
    if abs(len(headers1) - len(headers2)) > 1:
        return False

    norm1 = [_normalize_text(x) for x in headers1 if _normalize_text(x)]
    norm2 = [_normalize_text(x) for x in headers2 if _normalize_text(x)]

    if not norm1 or not norm2:
        return False

    if norm1 == norm2:
        return True

    set1 = set(norm1)
    matches = sum(1 for b in norm2 if b in set1)
    similarity = matches / max(len(norm1), len(norm2))
    return similarity >= 0.6


def _looks_like_header(row: List[str]) -> bool:
    if not row:
        return False
    non_empty = [_clean_cell(x) for x in row if _clean_cell(x)]
    if not non_empty:
        return False

    header_keywords = {
        "year", "date", "description", "name", "total", "quantity",
        "production", "capacity", "unit", "value", "parameter",
        "remarks", "location", "depth", "seam", "area", "coal", "sl", "no", "item"
    }

    matches = sum(1 for cell in non_empty if _normalize_text(cell) in header_keywords)
    if matches > 0:
        return True

    text_cells = sum(1 for cell in non_empty if re.search(r"[A-Za-z]", cell))
    return (text_cells / len(non_empty)) >= 0.6


# -----------------------------------------------------------------------------
# Table Extraction & Continuity Analysis (PyMuPDF + DOCX)
# -----------------------------------------------------------------------------

def _detect_pdf_table_continuations(pdf_path: str) -> List[Dict[str, Any]]:
    """
    Analyzes the source PDF with PyMuPDF to identify tables that span page boundaries.
    Returns a list of table metadata dicts with continuation flags.
    """
    if not HAS_PYMUPDF:
        return []

    doc = pymupdf.open(pdf_path)
    all_tables = []
    global_table_idx = 0

    for page_idx, page in enumerate(doc):
        page_num = page_idx + 1
        page_height = page.rect.height

        try:
            found_tables = page.find_tables()
        except Exception:
            found_tables = None

        if not found_tables or not found_tables.tables:
            continue

        for tbl_on_page_idx, tbl in enumerate(found_tables.tables):
            try:
                raw_rows = tbl.extract()
            except Exception:
                continue

            if not raw_rows:
                continue

            cleaned_rows = []
            for r in raw_rows:
                row_cells = [_clean_cell(c) for c in r]
                if any(row_cells):
                    cleaned_rows.append(row_cells)

            if not cleaned_rows:
                continue

            bbox = getattr(tbl, "bbox", (0, 0, 0, 0))
            top_y = bbox[1] if bbox else 0
            bottom_y = bbox[3] if bbox else page_height

            table_meta = {
                "global_idx": global_table_idx,
                "page_num": page_num,
                "tbl_on_page_idx": tbl_on_page_idx,
                "headers": cleaned_rows[0],
                "rows": cleaned_rows,
                "col_count": len(cleaned_rows[0]),
                "top_y": top_y,
                "bottom_y": bottom_y,
                "page_height": page_height,
                "is_continuation": False,
                "drop_repeated_header": False,
                "merges_into_idx": None
            }

            if global_table_idx > 0:
                prev_tbl = all_tables[-1]

                # Rule 1: Must be the first table on the page
                if tbl_on_page_idx == 0:
                    # Rule 2: Previous table finished near page bottom (>60% height)
                    prev_at_bottom = prev_tbl["bottom_y"] > (prev_tbl["page_height"] * 0.60)
                    # Rule 3: Current table starts near top of page (<45% height)
                    curr_at_top = top_y < (page_height * 0.45)
                    # Rule 4: Column count matches +/- 1
                    col_match = abs(len(cleaned_rows[0]) - prev_tbl["col_count"]) <= 1

                    if prev_at_bottom and curr_at_top and col_match:
                        first_row = cleaned_rows[0]
                        prev_headers = prev_tbl["headers"]

                        # Check if first row is a repeated header or pure data
                        if _headers_similar(prev_headers, first_row):
                            table_meta["is_continuation"] = True
                            table_meta["drop_repeated_header"] = True
                            table_meta["merges_into_idx"] = prev_tbl.get("merges_into_idx") or prev_tbl["global_idx"]
                        elif not _looks_like_header(first_row):
                            table_meta["is_continuation"] = True
                            table_meta["drop_repeated_header"] = False
                            table_meta["merges_into_idx"] = prev_tbl.get("merges_into_idx") or prev_tbl["global_idx"]

            all_tables.append(table_meta)
            global_table_idx += 1

    doc.close()
    return all_tables


def _are_tables_adjacent_in_docx(tbl_a, tbl_b) -> bool:
    """
    Checks if two tables in python-docx have no non-empty text paragraphs separating them.
    """
    curr = tbl_a._element.getnext()
    while curr is not None and curr != tbl_b._element:
        if curr.tag.endswith("p"):
            text = "".join(curr.itertext()).strip()
            if text:
                return False
        curr = curr.getnext()
    return True


def _join_docx_tables(doc: docx.Document, pdf_path: Optional[str] = None) -> int:
    """
    Identifies and merges continuous multi-page tables in python-docx at the XML level.
    Returns the count of table joins performed.
    """
    if len(doc.tables) < 2:
        return 0

    joins_count = 0

    # Strategy 1: Use PyMuPDF PDF analysis if pdf_path is provided
    pdf_tables = _detect_pdf_table_continuations(pdf_path) if (pdf_path and HAS_PYMUPDF) else []

    if pdf_tables and len(pdf_tables) == len(doc.tables):
        # Map 1:1 between PyMuPDF tables and doc.tables (Iterate backwards to safely remove XML elements)
        i = len(doc.tables) - 1
        while i > 0:
            meta = pdf_tables[i]
            if meta["is_continuation"] and meta["merges_into_idx"] is not None:
                target_idx = meta["merges_into_idx"]
                tbl_target = doc.tables[target_idx]
                tbl_curr = doc.tables[i]

                rows_to_move = tbl_curr.rows[1:] if meta["drop_repeated_header"] else tbl_curr.rows
                for row in rows_to_move:
                    tbl_target._tbl.append(row._tr)

                # Remove merged table from XML DOM
                tbl_curr._element.getparent().remove(tbl_curr._element)
                joins_count += 1
            i -= 1

    else:
        # Strategy 2: Direct DOCX structural inspection fallback
        i = 0
        while i < len(doc.tables) - 1:
            tbl_a = doc.tables[i]
            tbl_b = doc.tables[i + 1]

            cols_a = len(tbl_a.columns)
            cols_b = len(tbl_b.columns)

            if abs(cols_a - cols_b) <= 1 and _are_tables_adjacent_in_docx(tbl_a, tbl_b):
                headers_a = [c.text for c in tbl_a.rows[0].cells]
                first_row_b = [c.text for c in tbl_b.rows[0].cells]

                is_same = False
                drop_header = False

                if _headers_similar(headers_a, first_row_b):
                    is_same = True
                    drop_header = True
                elif not _looks_like_header(first_row_b):
                    is_same = True
                    drop_header = False

                if is_same:
                    rows_to_move = tbl_b.rows[1:] if drop_header else tbl_b.rows
                    for row in rows_to_move:
                        tbl_a._tbl.append(row._tr)

                    tbl_b._element.getparent().remove(tbl_b._element)
                    joins_count += 1
                    # Do not increment i so tbl_a can absorb subsequent continuations (e.g. Page 3, 4)
                    continue

            i += 1

    return joins_count


# -----------------------------------------------------------------------------
# Run Info & Heading Normalization Helpers
# -----------------------------------------------------------------------------

def _run_info(run) -> RunInfo:
    color_rgb = None
    try:
        if run.font.color and run.font.color.rgb:
            color_rgb = str(run.font.color.rgb)
    except Exception:
        pass
    return RunInfo(
        text=run.text,
        bold=bool(run.bold),
        italic=bool(run.italic),
        size_pt=run.font.size.pt if run.font.size else None,
        font_name=run.font.name,
        color_rgb=color_rgb,
    )


def _has_break(run) -> bool:
    return run._element.find(qn("w:br")) is not None


def _split_paragraph_runs(paragraph: Paragraph) -> list[list[RunInfo]]:
    """Split a paragraph's runs into segments wherever a <w:br/> run occurs."""
    segments: list[list[RunInfo]] = [[]]
    for run in paragraph.runs:
        if _has_break(run):
            segments.append([])
            continue
        if run.text == "":
            continue
        segments[-1].append(_run_info(run))
    return [seg for seg in segments if seg]


def _body_size_and_heading_levels(doc) -> tuple[float, dict[float, int]]:
    body_sizes = Counter()
    bold_sizes = Counter()
    for p in doc.paragraphs:
        for run in p.runs:
            if not run.text.strip():
                continue
            size = run.font.size.pt if run.font.size else None
            if size is None:
                continue
            if run.bold:
                bold_sizes[size] += 1
            else:
                body_sizes[size] += 1

    body_size = body_sizes.most_common(1)[0][0] if body_sizes else (
        bold_sizes.most_common(1)[0][0] if bold_sizes else 11.0
    )

    heading_sizes = sorted(
        {s for s in bold_sizes if s > body_size * SIZE_MARGIN},
        reverse=True,
    )[:MAX_HEADING_LEVEL]

    level_map = {size: i + 1 for i, size in enumerate(heading_sizes)}
    return body_size, level_map


def _segment_style(segment: list[RunInfo], body_size: float, level_map: dict[float, int]) -> str:
    first = segment[0]
    if first.bold and first.size_pt and first.size_pt in level_map:
        return f"Heading {level_map[first.size_pt]}"
    return "Normal"


def _write_segment(new_para: Paragraph, segment: list[RunInfo]):
    for info in segment:
        run = new_para.add_run(info.text)
        run.bold = info.bold
        run.italic = info.italic

        if info.size_pt:
            from docx.shared import Pt
            run.font.size = Pt(info.size_pt)
        if info.font_name:
            run.font.name = info.font_name
        if info.color_rgb:
            try:
                from docx.shared import RGBColor
                run.font.color.rgb = RGBColor.from_string(info.color_rgb)
            except Exception:
                pass


# -----------------------------------------------------------------------------
# Main Public Function
# -----------------------------------------------------------------------------

def normalize_headings(input_path: str, output_path: str, pdf_path: Optional[str] = None) -> dict:
    """
    1. Detects and joins multi-page continuous tables in python-docx (guided by PyMuPDF if pdf_path is provided).
    2. Strips floating drawings and inline images.
    3. Re-tags heading paragraphs with built-in Word Heading styles and splits
       merged heading+body paragraphs.
    
    Returns a status report dict.
    """
    doc = docx.Document(input_path)

    # Step 1: Detect and join continuous tables first
    tables_joined = _join_docx_tables(doc, pdf_path=pdf_path)

    # Step 2: Remove inline images & floating shapes
    for shape in list(doc.inline_shapes):
        try:
            shape._inline.getparent().remove(shape._inline)
        except Exception:
            pass

    body_size, level_map = _body_size_and_heading_levels(doc)

    paragraphs_split = 0
    headings_tagged = 0

    # Step 3: Process paragraphs, split break runs, and assign Heading styles
    for paragraph in list(doc.paragraphs):
        for drawing in paragraph._element.xpath(".//w:drawing"):
            drawing.getparent().remove(drawing)

        segments = _split_paragraph_runs(paragraph)
        if not segments:
            continue

        if len(segments) == 1:
            style = _segment_style(segments[0], body_size, level_map)
            if style != "Normal":
                paragraph.style = doc.styles[style]
                headings_tagged += 1
            continue

        # Multiple segments: rebuild as separate paragraphs before the original paragraph
        for segment in segments:
            style = _segment_style(segment, body_size, level_map)
            new_para = paragraph.insert_paragraph_before("", style=style if style != "Normal" else None)
            _write_segment(new_para, segment)
            if style != "Normal":
                headings_tagged += 1

        paragraph._element.getparent().remove(paragraph._element)
        paragraphs_split += 1

    doc.save(output_path)
    return {
        "tables_joined": tables_joined,
        "body_size_pt": body_size,
        "heading_levels_detected": level_map,
        "paragraphs_split": paragraphs_split,
        "headings_tagged": headings_tagged,
    }