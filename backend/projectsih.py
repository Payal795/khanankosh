import json
import re
from typing import Dict, List, Any, Optional

import numpy as np
import docx
from docx.enum.text import WD_BREAK
from docx.text.paragraph import Paragraph
from docx.table import Table

from langchain_chroma import Chroma
from database import TableDatabase
from model_clients import create_chat_model, create_embeddings


# =====================================================================
# 1. INITIALIZE LLM, VECTOR STORE & SCORER
# =====================================================================

embeddings = create_embeddings()

vectorstore = Chroma(
    collection_name="coal_ministry_docs",
    embedding_function=embeddings,
    persist_directory="./chroma_db"
)

llm = create_chat_model()


class RetrievalQualityScorer:
    """
    Evaluates retrieved RAG chunks to prevent hallucination
    from weak matches.
    """

    def __init__(self, minimum_relevance: float = 0.60):
        self.encoder = create_embeddings()
        self.min_relevance = minimum_relevance

    def evaluate_retrieved_context(
        self,
        user_query: str,
        retrieved_chunks: list
    ) -> tuple:

        if not retrieved_chunks:
            return False, 0.0

        query_vec = np.asarray(self.encoder.embed_query(user_query))
        chunk_vecs = np.asarray(self.encoder.embed_documents(retrieved_chunks))

        scores = np.dot(chunk_vecs, query_vec) / (
            np.linalg.norm(chunk_vecs, axis=1)
            * np.linalg.norm(query_vec)
        )

        top_score = float(np.max(scores))

        is_sufficient = top_score >= self.min_relevance

        return is_sufficient, round(top_score, 4)


# Instantiate scorer
scorer = RetrievalQualityScorer(
    minimum_relevance=0.60
)


# =====================================================================
# 2. DOCX BLOCK ITERATOR
# =====================================================================

def iter_block_items(parent):
    """
    Iterates through paragraphs and tables in their original
    document order.
    """

    if isinstance(parent, docx.document.Document):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._tc

    for child in parent_elm.iterchildren():

        if isinstance(
            child,
            docx.oxml.text.paragraph.CT_P
        ):
            yield Paragraph(child, parent)

        elif isinstance(
            child,
            docx.oxml.table.CT_Tbl
        ):
            yield Table(child, parent)


# =====================================================================
# 3. RAG RETRIEVAL
# =====================================================================

def retrieve_rag_context(
    query_text: str,
    top_k: int = 4
) -> str:

    """
    Retrieves semantic chunks from ChromaDB and evaluates
    their relevance before passing them to the LLM.
    """

    try:

        docs = vectorstore.similarity_search(
            query_text,
            k=top_k
        )

        if not docs:
            return "No relevant narrative context found."

        chunks = [
            doc.page_content
            for doc in docs
        ]

        is_sufficient, score = (
            scorer.evaluate_retrieved_context(
                query_text,
                chunks
            )
        )

        if not is_sufficient:

            print(
                f"[RAG] Rejected weak context for "
                f"'{query_text[:40]}...' "
                f"(Score: {score} < "
                f"{scorer.min_relevance})"
            )

            return (
                "No highly relevant narrative context found "
                f"(Confidence: {score})."
            )

        return "\n---\n".join(chunks)

    except Exception as e:

        print(
            f"RAG Retrieval Error: {e}"
        )

        return (
            "No relevant narrative context retrieved "
            "due to query error."
        )


# =====================================================================
# 4. SQL RETRIEVAL
# =====================================================================

def retrieve_sql_context(
    headers_str: str,
    user_topic: str
) -> str:

    """
    Dynamically generates SQL to retrieve exact numerical
    values from the SQLite table database.
    """

    db = None

    try:

        db = TableDatabase("./tables.db")

        schema = db.describe_schema()

        sql_prompt = f"""
You are an expert SQLite data analyst.

Based on this schema:

{schema}

Write a SQLite query to fetch data relevant to this topic:

"{user_topic}"

The query should match these column headers:

"{headers_str}"

Rules:

1. Always include `_is_total = 0` when using aggregates
   such as SUM or AVG.
2. Return ONLY the raw SQL query.
3. Do not use markdown.
4. Do not use ```sql.
5. Do not provide explanations.
"""

        response = llm.invoke(sql_prompt)

        sql_query = response.content.strip()

        sql_query = (
            sql_query
            .replace("```sql", "")
            .replace("```", "")
            .strip()
        )

        print(
            f"[SQL] Generated query:\n{sql_query}"
        )

        cols, rows = db.query(sql_query)

        if not rows:
            return (
                "No exact SQL data found "
                "for these metrics."
            )

        result = [
            f"Columns: {cols}"
        ]

        for row in rows[:15]:
            result.append(str(row))

        return "\n".join(result)

    except Exception as e:

        print(
            f"[SQL] Extraction failed: {e}"
        )

        return (
            f"SQL Extraction Failed: {e}"
        )

    finally:

        if db is not None:
            try:
                db.close()
            except Exception:
                pass


# =====================================================================
# 5. PARAGRAPH REWRITING
# =====================================================================

def rewrite_paragraph_with_llm(
    text: str,
    context: Dict[str, Optional[str]],
    user_topic: str,
    user_description: str
) -> str:

    active_heading = (
        context.get("heading4")
        or context.get("heading3")
        or context.get("heading2")
        or context.get("heading1")
        or ""
    )

    search_query = (
        f"{active_heading} "
        f"{user_topic} "
        f"{user_description}"
    )

    rag_context = retrieve_rag_context(
        search_query,
        top_k=3
    )

    prompt = f"""
You are a technical document writer.

Rewrite the paragraph below so that it incorporates
relevant details from the RAG Context while satisfying
the User Requirements.

==================================================
CONTEXT HIERARCHY
==================================================

Heading 1:
{context.get('heading1') or 'N/A'}

Heading 2:
{context.get('heading2') or 'N/A'}

Heading 3:
{context.get('heading3') or 'N/A'}

Heading 4:
{context.get('heading4') or 'N/A'}

==================================================
USER REQUIREMENTS
==================================================

Topic:
{user_topic}

Description:
{user_description}

==================================================
RETRIEVED RAG CONTEXT
==================================================

{rag_context}

==================================================
ORIGINAL TEXT
==================================================

{text}

==================================================
INSTRUCTIONS
==================================================

1. Rewrite the content maintaining a professional tone.
2. Incorporate specific factual details from the RAG Context.
3. Keep the length concise and aligned with the original context.
4. Do not invent factual information.
5. Output ONLY the rewritten paragraph.
6. Do not include meta-commentary.
7. Do not introduce the answer with phrases such as
   "Here is the rewritten paragraph".
"""

    try:

        response = llm.invoke(prompt)

        return response.content.strip()

    except Exception as e:

        print(
            f"[PARAGRAPH] LLM rewrite failed: {e}"
        )

        # Preserve original text instead of destroying
        # document content if the LLM call fails.
        return text


# =====================================================================
# 6. TABLE EXTRACTION
# =====================================================================

def extract_table_data(
    table: Table
) -> Dict[str, Any]:

    """
    Extracts a DOCX table into:

    {
        "headers": [...],
        "rows": [...]
    }

    The function uses the actual dimensions of the
    Word table instead of assuming fixed dimensions.
    """

    if table is None or not table.rows:

        return {
            "headers": [],
            "rows": []
        }

    headers = [
        cell.text.strip()
        for cell in table.rows[0].cells
    ]

    rows = []

    for row in table.rows[1:]:

        rows.append([
            cell.text.strip()
            for cell in row.cells
        ])

    return {
        "headers": headers,
        "rows": rows
    }


# =====================================================================
# 7. NORMALIZE LLM TABLE OUTPUT
# =====================================================================

def normalize_table_data(
    table_data: Dict[str, Any]
) -> Optional[Dict[str, Any]]:

    """
    Ensures that LLM-generated table data is rectangular.

    Every row will have exactly the same number of values
    as the headers.
    """

    if not isinstance(table_data, dict):

        print(
            "[TABLE] LLM returned non-dictionary data."
        )

        return None

    headers = table_data.get(
        "headers",
        []
    )

    rows = table_data.get(
        "rows",
        []
    )

    if not isinstance(headers, list):

        print(
            "[TABLE] Invalid headers format."
        )

        return None

    if not headers:

        return None

    normalized_headers = []

    for header in headers:

        if header is None:
            normalized_headers.append("N/A")

        else:
            normalized_headers.append(
                str(header).strip()
            )

    num_columns = len(
        normalized_headers
    )

    normalized_rows = []

    if not isinstance(rows, list):

        rows = []

    for row in rows:

        if not isinstance(row, (list, tuple)):

            row = [row]

        row = list(row)

        if len(row) < num_columns:

            row.extend(
                ["N/A"]
                * (num_columns - len(row))
            )

        elif len(row) > num_columns:

            row = row[:num_columns]

        cleaned_row = []

        for value in row:

            if value is None:

                cleaned_row.append("N/A")

            elif str(value).strip() == "":

                cleaned_row.append("N/A")

            else:

                cleaned_row.append(
                    str(value)
                )

        normalized_rows.append(
            cleaned_row
        )

    return {
        "headers": normalized_headers,
        "rows": normalized_rows
    }


# =====================================================================
# 8. TABLE REWRITING
# =====================================================================

def rewrite_table_with_llm(
    table: Table,
    context: Dict[str, Optional[str]],
    user_topic: str,
    user_description: str
) -> Optional[Dict[str, Any]]:

    """
    Rewrites table values using:

    1. ChromaDB semantic retrieval
    2. SQLite exact numerical retrieval
    3. LLM-based table restructuring

    The returned table is always normalized into
    rectangular JSON.
    """

    table_data = extract_table_data(
        table
    )

    if not table_data["headers"]:

        print(
            "[TABLE] Skipping table with no headers."
        )

        return None

    active_heading = (
        context.get("heading4")
        or context.get("heading3")
        or context.get("heading2")
        or context.get("heading1")
        or ""
    )

    headers_str = " ".join(
        table_data["headers"]
    )

    # ---------------------------------------------------------------
    # DUAL RETRIEVAL
    # ---------------------------------------------------------------

    search_query = (
        f"{active_heading} "
        f"{headers_str} "
        f"{user_topic}"
    )

    print(
        f"[TABLE] RAG query: {search_query}"
    )

    rag_context = retrieve_rag_context(
        search_query,
        top_k=4
    )

    sql_context = retrieve_sql_context(
        headers_str,
        user_topic
    )

    # ---------------------------------------------------------------
    # LLM PROMPT
    # ---------------------------------------------------------------

    prompt = f"""
You are a data analyst updating a technical report table.

Your task is to update the ORIGINAL TABLE using ONLY
information supported by the RETRIEVED CONTEXTS and
USER REQUIREMENTS.

==================================================
SECTION CONTEXT
==================================================

Heading 1:
{context.get('heading1') or 'N/A'}

Heading 2:
{context.get('heading2') or 'N/A'}

Heading 3:
{context.get('heading3') or 'N/A'}

Heading 4:
{context.get('heading4') or 'N/A'}

==================================================
USER REQUIREMENTS
==================================================

Topic:
{user_topic}

Description:
{user_description}

==================================================
RETRIEVED NARRATIVE CONTEXT
==================================================

{rag_context}

==================================================
EXACT DATABASE NUMBERS
==================================================

{sql_context}

==================================================
ORIGINAL TABLE
==================================================

{json.dumps(table_data, indent=2)}

==================================================
RULES
==================================================

1. Preserve the existing table's logical structure whenever possible.

2. Do NOT assume that the original table has a fixed
   number of columns.

3. Do NOT assume that the original table has a fixed
   number of rows.

4. If the retrieved context contains a NEW YEAR that is
   explicitly supported and belongs in the table, you may
   add that year as a new column.

5. If the retrieved context contains a NEW parameter or
   metric that clearly belongs in the table, you may add
   it as a new row.

6. Preserve existing years and columns unless the retrieved
   information explicitly indicates that they should change.

7. Prioritize EXACT DATABASE NUMBERS over narrative context
   when updating numerical values.

8. If a value cannot be found in the provided contexts,
   use "N/A".

9. NEVER invent, estimate, infer, or hallucinate numerical
   values.

10. Every row MUST contain exactly the same number of values
    as the headers.

11. If a newly added column has no data for a row,
    use "N/A".

12. If a newly added row has no data for a column,
    use "N/A".

13. Keep the returned table rectangular.

14. Return ONLY valid JSON.

15. Do NOT return markdown.

16. Do NOT return ```json.

17. Do NOT provide explanations.

==================================================
OUTPUT FORMAT
==================================================

{{
    "headers": [
        "Header1",
        "Header2",
        "Header3"
    ],
    "rows": [
        ["Value1", "Value2", "Value3"],
        ["Value1", "Value2", "Value3"]
    ]
}}
"""

    try:

        response = llm.invoke(
            prompt
        )

        raw_content = response.content.strip()

    except Exception as e:

        print(
            f"[TABLE] LLM generation failed: {e}"
        )

        return None

    # ---------------------------------------------------------------
    # CLEAN MARKDOWN WRAPPERS
    # ---------------------------------------------------------------

    raw_content = re.sub(
        r"^```(?:json)?\s*",
        "",
        raw_content,
        flags=re.IGNORECASE
    )

    raw_content = re.sub(
        r"\s*```$",
        "",
        raw_content
    )

    raw_content = raw_content.strip()

    # ---------------------------------------------------------------
    # PARSE JSON
    # ---------------------------------------------------------------

    try:

        updated_table_data = json.loads(
            raw_content
        )

    except json.JSONDecodeError as e:

        print(
            "[TABLE] Failed to parse LLM JSON output."
        )

        print(
            f"[TABLE] JSON error: {e}"
        )

        print(
            "[TABLE] Raw LLM output:"
        )

        print(
            raw_content
        )

        return None

    # ---------------------------------------------------------------
    # NORMALIZE TABLE
    # ---------------------------------------------------------------

    updated_table_data = normalize_table_data(
        updated_table_data
    )

    if updated_table_data is None:

        print(
            "[TABLE] Invalid normalized table."
        )

        return None

    print(
        "[TABLE] LLM table generated successfully."
    )

    print(
        f"[TABLE] Generated columns: "
        f"{len(updated_table_data['headers'])}"
    )

    print(
        f"[TABLE] Generated rows: "
        f"{len(updated_table_data['rows'])}"
    )

    return updated_table_data


# =====================================================================
# 9. SAFE DOCX TABLE UPDATE
# =====================================================================

def update_docx_table(
    table: Table,
    new_data: Dict[str, Any]
):
    """
    Safely updates an existing DOCX table.

    We do NOT use row._tr.add_tc() because python-docx
    may not expose newly-created XML cells consistently
    through row.cells.

    Instead, the function works with the cells that
    actually exist in the DOCX table.
    """

    if table is None:

        print(
            "[TABLE UPDATE] Table is None."
        )

        return

    if not table.rows:

        print(
            "[TABLE UPDATE] Table contains no rows."
        )

        return

    if not isinstance(new_data, dict):

        print(
            "[TABLE UPDATE] Invalid table data."
        )

        return

    headers = new_data.get(
        "headers",
        []
    )

    rows = new_data.get(
        "rows",
        []
    )

    if not isinstance(headers, list):

        print(
            "[TABLE UPDATE] Headers are not a list."
        )

        return

    if not headers:

        print(
            "[TABLE UPDATE] No headers returned."
        )

        return

    if not isinstance(rows, list):

        rows = []

    # ---------------------------------------------------------------
    # ACTUAL DOCX DIMENSIONS
    # ---------------------------------------------------------------

    actual_docx_rows = len(
        table.rows
    )

    actual_docx_columns = max(
        (
            len(row.cells)
            for row in table.rows
        ),
        default=0
    )

    generated_columns = len(
        headers
    )

    generated_rows = len(
        rows
    )

    print(
        "\n[TABLE UPDATE]"
    )

    print(
        f"  DOCX rows    : {actual_docx_rows}"
    )

    print(
        f"  DOCX columns : {actual_docx_columns}"
    )

    print(
        f"  LLM rows     : {generated_rows + 1}"
    )

    print(
        f"  LLM columns  : {generated_columns}"
    )

    # ---------------------------------------------------------------
    # DIMENSION CHECK
    # ---------------------------------------------------------------

    if generated_columns > actual_docx_columns:

        print(
            "[TABLE UPDATE WARNING] "
            f"LLM generated {generated_columns} columns, "
            f"but DOCX contains only "
            f"{actual_docx_columns} columns."
        )

        print(
            "[TABLE UPDATE WARNING] "
            "Extra generated columns will be ignored "
            "to prevent DOCX corruption."
        )

    if generated_rows + 1 > actual_docx_rows:

        print(
            "[TABLE UPDATE WARNING] "
            f"LLM generated {generated_rows + 1} rows, "
            f"but DOCX contains only "
            f"{actual_docx_rows} rows."
        )

        print(
            "[TABLE UPDATE WARNING] "
            "Extra generated rows will be ignored."
        )

    # ---------------------------------------------------------------
    # UPDATE HEADER
    # ---------------------------------------------------------------

    header_cells = table.rows[0].cells

    header_columns_to_write = min(
        len(header_cells),
        generated_columns
    )

    for col_idx in range(
        header_columns_to_write
    ):

        value = headers[col_idx]

        if value is None:

            value = "N/A"

        else:

            value = str(value).strip()

            if not value:

                value = "N/A"

        try:

            header_cells[col_idx].text = value

        except IndexError:

            print(
                "[TABLE UPDATE WARNING] "
                f"Header column {col_idx} "
                "does not exist."
            )

    # ---------------------------------------------------------------
    # CLEAR UNUSED HEADER CELLS
    # ---------------------------------------------------------------

    if generated_columns < len(header_cells):

        for col_idx in range(
            generated_columns,
            len(header_cells)
        ):

            try:

                header_cells[col_idx].text = "N/A"

            except IndexError:

                continue

    # ---------------------------------------------------------------
    # UPDATE DATA ROWS
    # ---------------------------------------------------------------

    for row_idx, row_data in enumerate(rows):

        # +1 because row 0 is the header
        docx_row_idx = row_idx + 1

        if docx_row_idx >= len(
            table.rows
        ):

            print(
                "[TABLE UPDATE WARNING] "
                f"Generated row {row_idx} "
                "has no corresponding DOCX row."
            )

            break

        docx_row = table.rows[
            docx_row_idx
        ]

        cells = docx_row.cells

        if not isinstance(
            row_data,
            (list, tuple)
        ):

            row_data = [row_data]

        row_data = list(
            row_data
        )

        # -----------------------------------------------------------
        # NORMALIZE ROW LENGTH
        # -----------------------------------------------------------

        if len(row_data) < generated_columns:

            row_data.extend(
                ["N/A"]
                * (
                    generated_columns
                    - len(row_data)
                )
            )

        elif len(row_data) > generated_columns:

            row_data = row_data[
                :generated_columns
            ]

        # -----------------------------------------------------------
        # WRITE EXISTING CELLS ONLY
        # -----------------------------------------------------------

        columns_to_write = min(
            len(cells),
            generated_columns,
            len(row_data)
        )

        for col_idx in range(
            columns_to_write
        ):

            value = row_data[
                col_idx
            ]

            if value is None:

                value = "N/A"

            else:

                value = str(value).strip()

                if not value:

                    value = "N/A"

            try:

                cells[col_idx].text = value

            except IndexError:

                print(
                    "[TABLE UPDATE WARNING] "
                    f"Could not access "
                    f"row={docx_row_idx}, "
                    f"column={col_idx}."
                )

                continue

        # -----------------------------------------------------------
        # CLEAR REMAINING EXISTING CELLS
        # -----------------------------------------------------------

        if len(cells) > columns_to_write:

            for col_idx in range(
                columns_to_write,
                len(cells)
            ):

                try:

                    cells[col_idx].text = "N/A"

                except IndexError:

                    continue

        # -----------------------------------------------------------
        # WARN ABOUT EXTRA GENERATED COLUMNS
        # -----------------------------------------------------------

        if len(row_data) > len(cells):

            print(
                "[TABLE UPDATE WARNING] "
                f"Row {row_idx}: generated "
                f"{len(row_data)} values but DOCX "
                f"contains only {len(cells)} cells. "
                "Extra values ignored."
            )

    # ---------------------------------------------------------------
    # FINAL EMPTY CELL CLEANUP
    # ---------------------------------------------------------------

    for row in table.rows:

        cells = row.cells

        for cell in cells:

            try:

                if not cell.text.strip():

                    cell.text = "N/A"

            except Exception:

                continue

    print(
        "[TABLE UPDATE] Completed safely."
    )


# =====================================================================
# 10. REMOVE EMPTY PARAGRAPHS
# =====================================================================

def remove_empty_paragraphs(doc):
    """
    Removes completely empty paragraphs from the document.

    Paragraphs containing page breaks are preserved.
    This prevents large unwanted white spaces caused by
    leftover blank paragraphs.
    """

    removed_count = 0

    # Work through a copy because we are modifying
    # the document while iterating.
    for paragraph in list(doc.paragraphs):

        # -----------------------------------------------------------
        # CHECK FOR PAGE BREAK
        # -----------------------------------------------------------

        has_page_break = False

        for run in paragraph.runs:

            try:

                page_breaks = run._r.xpath(
                    ".//w:br[@w:type='page']"
                )

                if page_breaks:

                    has_page_break = True
                    break

            except Exception:

                continue

        # Never remove a paragraph containing a page break.
        if has_page_break:
            continue

        # -----------------------------------------------------------
        # REMOVE EMPTY PARAGRAPH
        # -----------------------------------------------------------

        if not paragraph.text.strip():

            try:

                p = paragraph._element

                p.getparent().remove(p)

                removed_count += 1

            except Exception as e:

                print(
                    "[DOCUMENT CLEANUP WARNING] "
                    f"Could not remove empty paragraph: {e}"
                )

    print(
        f"[DOCUMENT CLEANUP] Removed "
        f"{removed_count} empty paragraphs."
    )


# =====================================================================
# 11. MAIN DOCUMENT PROCESSING PIPELINE
# =====================================================================

def process_and_rewrite_document(
    input_path: str,
    output_path: str,
    user_topic: str,
    user_description: str
):

    print(
        "\n=================================================="
    )

    print(
        "STARTING DOCUMENT REWRITE"
    )

    print(
        "=================================================="
    )

    print(
        f"Input : {input_path}"
    )

    print(
        f"Output: {output_path}"
    )

    print(
        f"Topic : {user_topic}"
    )

    print(
        f"Description: {user_description}"
    )

    print(
        "=================================================="
    )

    # ---------------------------------------------------------------
    # LOAD DOCX
    # ---------------------------------------------------------------

    try:

        doc = docx.Document(
            input_path
        )

    except Exception as e:

        print(
            f"[DOCUMENT] Failed to open DOCX: {e}"
        )

        raise

    # ---------------------------------------------------------------
    # HEADING STATE
    # ---------------------------------------------------------------

    heading1 = None
    heading2 = None
    heading3 = None
    heading4 = None

    first_heading1_seen = False

    # ---------------------------------------------------------------
    # PROCESS DOCUMENT BLOCKS
    # ---------------------------------------------------------------

    for block in iter_block_items(doc):

        # ===========================================================
        # PARAGRAPH
        # ===========================================================

        if isinstance(
            block,
            Paragraph
        ):

            style = block.style.name

            text = block.text.strip()

            if not text:

                continue

            # -------------------------------------------------------
            # HEADING 1
            # -------------------------------------------------------

            if style == "Heading 1":

                heading1 = text

                heading2 = None
                heading3 = None
                heading4 = None

                # ---------------------------------------------------
                # PAGE BREAK
                # ---------------------------------------------------
                #
                # IMPORTANT:
                # We add the page break directly to the Heading 1
                # instead of creating a separate blank paragraph.
                #
                # This prevents unnecessary white space.
                # ---------------------------------------------------

                if first_heading1_seen:

                    if block.runs:

                        block.runs[0].add_break(
                            WD_BREAK.PAGE
                        )

                    else:

                        block.add_run().add_break(
                            WD_BREAK.PAGE
                        )

                else:

                    first_heading1_seen = True

            # -------------------------------------------------------
            # HEADING 2
            # -------------------------------------------------------

            elif style == "Heading 2":

                heading2 = text

                heading3 = None
                heading4 = None

            # -------------------------------------------------------
            # HEADING 3
            # -------------------------------------------------------

            elif style == "Heading 3":

                heading3 = text

                heading4 = None

            # -------------------------------------------------------
            # HEADING 4
            # -------------------------------------------------------

            elif style == "Heading 4":

                heading4 = text

            # -------------------------------------------------------
            # NORMAL PARAGRAPH
            # -------------------------------------------------------

            else:

                context = {
                    "heading1": heading1,
                    "heading2": heading2,
                    "heading3": heading3,
                    "heading4": heading4
                }

                print(
                    "\n[PARAGRAPH]"
                )

                print(
                    f"Heading: "
                    f"{heading4 or heading3 or heading2 or heading1}"
                )

                try:

                    new_text = (
                        rewrite_paragraph_with_llm(
                            text=text,
                            context=context,
                            user_topic=user_topic,
                            user_description=user_description
                        )
                    )

                    block.text = new_text

                except Exception as e:

                    print(
                        "[PARAGRAPH ERROR]"
                    )

                    print(
                        f"{e}"
                    )

                    # Keep original paragraph
                    # instead of crashing the complete report.
                    continue

        # ===========================================================
        # TABLE
        # ===========================================================

        elif isinstance(
            block,
            Table
        ):

            context = {
                "heading1": heading1,
                "heading2": heading2,
                "heading3": heading3,
                "heading4": heading4
            }

            print(
                "\n=================================================="
            )

            print(
                "[TABLE]"
            )

            print(
                f"Section: "
                f"{heading4 or heading3 or heading2 or heading1}"
            )

            print(
                "=================================================="
            )

            try:

                updated_table_data = (
                    rewrite_table_with_llm(
                        table=block,
                        context=context,
                        user_topic=user_topic,
                        user_description=user_description
                    )
                )

                if updated_table_data:

                    update_docx_table(
                        block,
                        updated_table_data
                    )

                else:

                    print(
                        "[TABLE] No updated table returned. "
                        "Original table preserved."
                    )

            except Exception as e:

                print(
                    "[TABLE ERROR]"
                )

                print(
                    f"Exception: {type(e).__name__}"
                )

                print(
                    f"Message: {e}"
                )

                # IMPORTANT:
                # Do not crash the entire document because
                # one table failed.
                continue

    # ---------------------------------------------------------------
    # REMOVE EMPTY PARAGRAPHS
    # ---------------------------------------------------------------

    print(
        "\n=================================================="
    )

    print(
        "[DOCUMENT] CLEANING EMPTY PARAGRAPHS"
    )

    print(
        "=================================================="
    )

    remove_empty_paragraphs(
        doc
    )

    # ---------------------------------------------------------------
    # SAVE
    # ---------------------------------------------------------------

    try:

        doc.save(
            output_path
        )

    except Exception as e:

        print(
            f"[DOCUMENT] Failed to save DOCX: {e}"
        )

        raise

    print(
        "\n=================================================="
    )

    print(
        "DOCUMENT PROCESSING COMPLETE"
    )

    print(
        f"Successfully saved to: {output_path}"
    )

    print(
        "=================================================="
    )

    return output_path